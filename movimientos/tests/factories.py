"""Synthetic xlsx builders that reproduce the real bank layouts (made-up data only)."""
from pathlib import Path

import openpyxl

CARD_HEADER = ["Fecha", "Descripción", "Cuotas", "Comprobante", "Monto en pesos", "Monto en dólares"]

# Section rows for a synthetic Visa statement. ARS consumption nets to 20000.00 + 31041.94 - 31041.94.
DEFAULT_CARD_ROWS = [
    ["Pago de tarjeta y devoluciones"], CARD_HEADER,
    ["25/08/2026", "Su pago en pesos", "", "-", "$-10.000,00", None],
    [None, None, None, None, None, None],
    [None, None, None, None, None, None],
    ["Tarjeta de Juan Perez - Visa Crédito terminada en 5544"], CARD_HEADER,
    ["24/07/2025", "Tienda ejemplo uno ", "14 de 18", "111111*", "$31.041,94", None],
    ["20/08/2026", "Servicio ejemplo usd ", "", "222222K", None, "U$S20,00"],
    [None, "Servicio ejemplo dos ", "", "555555K", None, "U$S3,71"],
    [None, "Tienda ejemplo uno ", "", "333333K", "$-31.041,94", None],
    ["22/08/2026", "Comercio ejemplo ", "", "444444F", "$20.000,00", None],
    ["Total de Visa Crédito terminada en 5544", None, None, None, "$20.000,00", "U$S23,71"],
    [None, None, None, None, None, None],
    [None, None, None, None, None, None],
    ["Otros conceptos"], ["Descripción", "Monto en pesos"],
    ["Impuesto de sellos $", "$1.234,56"],
    [None, None],
    ["Aviso importante: archivo sin validez legal."],
]


def make_card_xlsx(tmpdir, rows_spec=None, card="Visa Crédito", last4="5544",
                   cierre="24/09/2026", vencimiento="05/10/2026") -> Path:
    rows = [
        [None], ["Movimientos del resumen"], [f"Tarjeta {card} terminada en {last4}"], [None],
        ["Fecha de cierre", "Fecha de vencimiento"], [cierre, vencimiento],
        ["Total a pagar"], ["$1,00", "U$S1,00"], [None], [None],
        ["Cierres y vencimientos"], [f"Cierre: {cierre}"], [None], [None],
    ]
    rows += DEFAULT_CARD_ROWS if rows_spec is None else rows_spec
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    for r in rows:
        ws.append(r)
    path = Path(tmpdir) / "resumen.xlsx"
    wb.save(path)
    return path


ACCOUNT_HEADER = [None, "Fecha", "Sucursal origen", "Descripción", "Referencia",
                  "Caja de Ahorro", "Cuenta Corriente", "Saldo"]

DEFAULT_ACCOUNT_ROWS = [
    ["19/10/2026", "000 - CASA CENTRAL", "Pago de servicios \t Edesur: 0001928991", "45645645", "-82212.63", "", ""],
    ["07/10/2026", "009 - SUCURSAL X", "Impuesto ley 25.413 debito 0,6% \t ", "000011824", "", "-8.92", "2754646.73"],
    ["13/10/2026", "000 - CASA CENTRAL", "Debito automatico \t Seguro ejemplo-0600001", "00406031", "-15337.00", "", ""],
    ["13/10/2026", "000 - CASA CENTRAL", "Debito automatico \t Seguro ejemplo-0600002", "00406030", "-15337.00", "", ""],
    ["13/10/2026", "000 - CASA CENTRAL", "Debito automatico \t Seguro ejemplo-0600003", "00406029", "-15337.00", "", ""],
    ["14/10/2026", "000 - CASA CENTRAL", "Transferencia recibida \t De ejemplo", "00003872", "2110004.00", "", "3611748.88"],
]


def make_account_xlsx(tmpdir, rows=None, moneda="Pesos") -> Path:
    data = DEFAULT_ACCOUNT_ROWS if rows is None else rows
    sheet = [
        [","], [None], [None, None, None, None, None, None, None, "lunes, 19 de octubre de 2026 - 19:7"],
        [None], [None], [None],
        [None, "Cuenta", "Cuenta única 000-000000/0"],
        [None, "Moneda", moneda],
        [None, "Fecha", "20/08/2026 - 19/10/2026"],
        [None], [None], [None, "Últimos movimientos"], [None], ACCOUNT_HEADER,
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "UltimosMovimientos"
    for r in sheet:
        ws.append(r)
    for r in data:
        ws.append([None] + list(r))
    for _ in range(3):
        ws.append([None])
    path = Path(tmpdir) / "movimientos.xlsx"
    wb.save(path)
    return path


# Galicia "Resumen de Cuenta Sueldo" PDF as extracted text lines (made-up data). Running balance:
# 1.000.000,00 -6.000 +500.000,50 -50.000 -50.000 +1.000 = 1.395.000,50.
GALICIA_PDF_LINES = [
    "Resumen de Cuenta Sueldo",
    "JUAN PEREZ CUIT del Responsable Impositivo : 20-00000000-1",
    "Datos de la cuenta Período de movimientos Saldos",
    "Tipo de cuenta",
    "Cuenta Sueldo",
    "$1.000.000,00",
    "Número de cuenta",
    "N° 1234567-8 999-1",
    "$1.395.000,50",
    "CBU 28/08/2026 25/09/2026",
    "0070999130001234567812",
    "Disponés de 30 días desde la recepción para cuestionar este resumen.",
    "Movimientos",
    "Fecha Descripción Origen Crédito Débito Saldo",
    "31/08/26 TRANSFERENCIA A TERCEROS -6.000,00 994.000,00",
    "Persona Ejemplo",
    "20000000001",
    "VARIOS",
    "31/08/26 ACREDITAMIENTO DE HABERES 500.000,50 1.494.000,50",
    "EMPRESA EJEMPLO",
    "Resumen de Cuenta Sueldo Página 1 / 2",
    "20260925000000000P",
    "Resumen de Cuenta Sueldo",
    "Fecha Descripción Origen Crédito Débito Saldo",
    "ACRED.HABERES",
    "11/09/26 TRANSF. CTAS PROPIAS -50.000,00 1.444.000,50",
    "JUAN PEREZ",
    "11/09/26 TRANSF. CTAS PROPIAS -50.000,00 1.394.000,50",
    "JUAN PEREZ",
    "16/09/26 INTERES CAPITALIZADO 1.000,00 1.395.000,50",
    "Septiembre 2026",
    "Total $ 501.000,50 -$ 106.000,00 $ 1.395.000,50",
    "Los depósitos en pesos y en moneda extranjera cuentan con la garantía de hasta $25.000.000.",
    "Canales de atención",
    "Banco de Galicia y Buenos Aires S.A.U.",
    "Resumen de Cuenta Sueldo Página 2 / 2",
    "20260925000000000P",
]

# The same movements as the Galicia xlsx lists them: newest first, multi-line text, balances
# without thousands separators.
GALICIA_XLSX_ROWS = [
    ["16/09/2026", " INTERES CAPITALIZADO\n Septiembre 2026\n", "0,00", "1.000,00", "1395000,50", None],
    ["11/09/2026", " TRANSF. CTAS PROPIAS\n JUAN  PEREZ\n", "-50.000,00", "0,00", "1394000,50", None],
    ["11/09/2026", " TRANSF. CTAS PROPIAS\n JUAN PEREZ\n", "-50.000,00", "0,00", "1444000,50", None],
    ["31/08/2026", " ACREDITAMIENTO DE HABERES\n EMPRESA EJEMPLO S.A.\n ACRED.HABERES\n", "0,00", "500.000,50",
     "1494000,50", None],
    ["31/08/2026", " TRANSFERENCIA A TERCEROS\n Persona Ejemplo\n 20000000001\n VARIOS\n", "-6.000,00", "0,00",
     "994000,00", None],
]
GALICIA_XLSX_HEADER = ["Fecha", "Movimiento", "Débito", "Crédito", "Saldo Parcial", "Comentarios"]


def make_galicia_xlsx(tmpdir, rows=None, moneda="Pesos") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cuentas"
    for r in ([f"Banco Galicia - Caja Ahorro {moneda}"], ["Nro. de Cuenta: ...9999991"],
              ["Fecha Actual: 22/10/2026"], ["Hora Actual: 00:49"],
              ["Intervalo de Consulta: del 22/08/2026 al 23/10/2026"], GALICIA_XLSX_HEADER):
        ws.append(r)
    for r in (GALICIA_XLSX_ROWS if rows is None else rows):
        ws.append(list(r))
    path = Path(tmpdir) / "Extracto_00000000000.xlsx"
    wb.save(path)
    return path
